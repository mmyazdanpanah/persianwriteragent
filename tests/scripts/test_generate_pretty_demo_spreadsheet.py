# WriterAgent - AI Writing Assistant for LibreOffice
# Copyright (c) 2026 KeithCu (modifications and relicensing)
#
# SPDX-License-Identifier: GPL-3.0-or-later
"""Unit tests for scripts/generate_pretty_demo_spreadsheet.py."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from scripts.generate_pretty_demo_spreadsheet import (
    CALC_PYTHON_ADDIN_FN,
    ENG_RANGE,
    FORECAST_RANGE_ODS_CROSS,
    MARKETING_NAMED_RANGE,
    MARKETING_RANGE_XLSX,
    RESULTS_PY_CODE_MAX_LEN,
    SALES_NAMED_RANGE,
    SALES_RANGE_ODS_CROSS,
    SALES_RANGE_XLSX,
    SALES_ZIPS_BY_REGION,
    SQL_IDENTITY_TEACH,
    SQL_RESULTS_SPILL_GUTTER_COLS,
    SQL_RESULTS_SPILL_GUTTER_ROWS,
    SQL_SALES_BY_REGION_CATEGORY,
    SQL_SALES_ZIP_INCOME_JOIN,
    STANDARD_METRICS_BANNER,
    ZIP_INCOME_CSV_NAME,
    ZIP_INCOME_FIXTURE,
    _ODS_SHEET_COLUMNS,
    _ods_named_range_address,
    _scenario_result_formula,
    _xlsx_named_range_address,
    build_ods_showcase,
    build_xlsx_showcase,
    duckdb_join_from_cell_code,
    duckdb_sql_from_cell_code,
    get_marketing_dataset,
    get_sales_dataset,
    ods_formula,
    py_formula,
    relative_named_range_eval_start_row,
    sql_demo_scenarios,
    sql_query_lines,
    sql_results_gutter_rows,
    standard_sheet_specs,
    write_zip_income_csv,
)

# Tokens that must live in sheet cells, never inside the RESULTS formula string.
_SQL_EMBED_MARKERS = (
    "SUM(Revenue)",
    "SUM(Ad_Spend)",
    "GROUP BY",
    "COUNT(*)",
    "NULLIF",
    "FROM sales",
    "FROM marketing",
)

FIXTURE_ODS = Path(__file__).resolve().parents[1] / "fixtures" / "python_showcase_demo.ods"
_NESTED_SHEET_REF = re.compile(r"\[\$[A-Za-z0-9_]*\[\$")


def _ods_content_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        return zf.read("content.xml").decode("utf-8")


def test_ods_formula_sales_analytics_range_not_rematched() -> None:
    out = ods_formula('=PY("x"; Sales_Analytics.A5:I40)')
    assert "[$Sales_Analytics.A5:.I40]" in out
    assert _NESTED_SHEET_REF.search(out) is None
    assert "[$S[$ales_Analytics" not in out
    assert out.startswith(f"of:={CALC_PYTHON_ADDIN_FN}(")


def test_ods_formula_forecasting_range_not_rematched() -> None:
    out = ods_formula('=PY("y"; Forecasting.A1:B2)')
    assert "[$Forecasting.A1:.B2]" in out
    assert _NESTED_SHEET_REF.search(out) is None
    assert "[$F[$orecasting" not in out


def test_ods_formula_same_sheet_range_and_cell() -> None:
    out = ods_formula('=PY("sum(r[7] for r in data[1:])"; A5:I40; F46)')
    assert "[.A5:.I40]" in out
    assert "[.F46]" in out
    assert "A5:I40" not in out
    assert out.startswith(f"of:={CALC_PYTHON_ADDIN_FN}(")


def test_ods_formula_python_wrapper_and_cross_sheet_cell() -> None:
    out = ods_formula('=PYTHON("x"; Sales_Analytics.F47)')
    assert out.startswith(f"of:={CALC_PYTHON_ADDIN_FN}(")
    assert "[$Sales_Analytics.F47]" in out
    assert _NESTED_SHEET_REF.search(out) is None


def test_ods_formula_leaves_non_formula_text() -> None:
    assert ods_formula("plain") == "plain"


def test_ods_engineering_temperature_uses_openformula_semicolon() -> None:
    """Regression: card 3 was ``=PY("...", A5:E11)`` (comma, old row-5 range)."""
    eng = next(s for s in standard_sheet_specs() if s["name"] == "Engineering_Math")
    title, unused_desc, code, args = next(m for m in eng["metrics"] if "Temperature" in m[0])
    assert title.startswith("3. Temperature")
    assert args == (ENG_RANGE,)
    assert ENG_RANGE == "A4:E10"
    formula = py_formula(code, *args, ods=True)
    assert formula == f'=PY("{code}"; {ENG_RANGE})'
    assert ", A" not in formula
    out = ods_formula(formula)
    assert "[.A4:.E10]" in out
    assert ", [.A" not in out
    assert "; [.A4:.E10]" in out or ";[.A4:.E10]" in out


def test_ods_formula_statistics_ml_range_not_rematched() -> None:
    out = ods_formula('=PY("x"; Statistics_ML.A5:G25)')
    assert "[$Statistics_ML.A5:.G25]" in out
    assert _NESTED_SHEET_REF.search(out) is None


def _quoted_formula_payload(formula: str) -> str:
    """First ``"…"`` argument of ``=FN("code", range)``. Exactly one pair of quotes."""
    assert formula.startswith("="), formula
    first = formula.index('"')
    last = formula.rindex('"')
    assert first < last, formula
    # A premature " (e.g. Python """) would add extra quotes and leak SQL
    # tokens like SUM(...) into the formula — Calc Err:508.
    assert formula.count('"') == 2, formula
    return formula[first + 1 : last]


def test_sales_dataset_has_zip_column_from_real_zctas() -> None:
    grid = get_sales_dataset()
    assert grid[0][-1] == "ZIP"
    assert grid[0].index("Revenue") == 7
    zips = {str(row[-1]) for row in grid[1:]}
    allowed = {z for group in SALES_ZIPS_BY_REGION.values() for z in group}
    assert zips <= allowed
    assert zips  # at least one assigned


def _is_sql_results_formula(formula: str) -> bool:
    """Live SQL_DuckDB RESULTS: sheet-only ``con.sql`` or join ``run_sql``."""
    return "con.sql(sql)" in formula or "run_sql(" in formula


def test_sql_demo_scenarios_include_sheet_and_join() -> None:
    kinds = {s["kind"] for s in sql_demo_scenarios()}
    assert kinds == {"sheet_sales", "sheet_marketing", "join_zip"}
    join = next(s for s in sql_demo_scenarios() if s["kind"] == "join_zip")
    assert "zip_income" in join["sql"]
    assert join["sql"] == SQL_SALES_ZIP_INCOME_JOIN
    assert "ZIP" in join["sql"] or "zip" in join["sql"]
    teach = " ".join(s["blurb"] for s in sql_demo_scenarios()) + SQL_IDENTITY_TEACH
    assert "{sheet" in teach
    assert "{named_range" in teach
    assert "file#Sheet" in teach or "budget.xlsx#Actuals" in teach
    assert SALES_NAMED_RANGE in teach


def test_zip_income_csv_is_fuller_acs_extract_and_covers_sales_zips() -> None:
    assert ZIP_INCOME_FIXTURE.is_file()
    lines = ZIP_INCOME_FIXTURE.read_text(encoding="utf-8").splitlines()
    assert lines[0].startswith("zip,median_household_income")
    # Fuller ZCTA extract (not a 40-row toy); ACS has ~33k ZCTAs.
    assert len(lines) > 5000
    csv_zips = {line.split(",", 1)[0] for line in lines[1:] if line.strip()}
    sales_zips = {z for group in SALES_ZIPS_BY_REGION.values() for z in group}
    missing = sales_zips - csv_zips
    assert not missing, f"sales ZIPs missing from ACS extract: {missing}"


def test_write_zip_income_csv_copies_sibling(tmp_path: Path) -> None:
    dest = write_zip_income_csv(tmp_path)
    assert dest.name == ZIP_INCOME_CSV_NAME
    assert dest.is_file()
    assert dest.stat().st_size == ZIP_INCOME_FIXTURE.stat().st_size


def test_generated_ods_has_sql_duckdb_sheet(tmp_path: Path) -> None:
    out = tmp_path / "python_showcase_demo.ods"
    build_ods_showcase(out)
    write_zip_income_csv(tmp_path)
    assert out.is_file()
    assert (tmp_path / ZIP_INCOME_CSV_NAME).is_file()

    from odf.opendocument import load
    from odf.table import Table
    from odf.text import P

    doc = load(str(out))
    names = [str(t.getAttribute("name")) for t in doc.spreadsheet.getElementsByType(Table)]
    assert "Sales_Analytics" in names
    assert "SQL_DuckDB" in names
    sql_sheet = next(t for t in doc.spreadsheet.getElementsByType(Table) if t.getAttribute("name") == "SQL_DuckDB")
    text = "\n".join(str(p) for p in sql_sheet.getElementsByType(P))
    assert "zip_income" in text
    assert "GROUP BY" in text or "Region" in text
    assert "{sheet" in text and "{named_range" in text
    assert SALES_NAMED_RANGE in text
    _assert_ods_results_unmerged_with_clearance(sql_sheet)


def _assert_ods_formulas_and_layout(xml: str) -> None:
    sales_of = f"[${SALES_RANGE_ODS_CROSS.replace(':', ':.')}]"
    assert sales_of in xml
    forecast_of = f"[${FORECAST_RANGE_ODS_CROSS.replace(':', ':.')}]"
    assert forecast_of in xml
    # Engineering °C→°F used a comma before the range; OpenFormula needs ``;``.
    assert (
        "9/5 + 32, 1)&quot;; [.A4:.E10]" in xml
        or '9/5 + 32, 1)"; [.A4:.E10]' in xml
    )
    assert "[$S[$ales_Analytics" not in xml
    assert "[$F[$orecasting" not in xml
    assert "[$S[$tatistics_ML" not in xml
    assert _NESTED_SHEET_REF.search(xml) is None
    assert xml.count("<table:table-column") >= sum(len(cols) for cols in _ODS_SHEET_COLUMNS.values())
    assert "style:column-width" in xml
    assert "style:row-height" in xml
    # Spans without covered placeholders collapse Overview KPIs in headed Calc.
    assert "<table:covered-table-cell" in xml
    assert "SQL_DuckDB" in xml
    # SQL text is cell content; RESULTS formulas are short OpenFormula runners.
    assert "SUM(Revenue)" in xml
    assert "SUM(Ad_Spend)" in xml
    assert SALES_NAMED_RANGE in xml
    assert MARKETING_NAMED_RANGE in xml
    formula_attrs = re.findall(r'table:formula="([^"]*)"', xml)
    # ``data[1]`` also appears in Forecasting CGR; pin RESULTS on the DuckDB runners.
    results = [f for f in formula_attrs if _is_sql_results_formula(f)]
    assert len(results) == 3
    sheet_only = [f for f in results if "con.sql(sql)" in f]
    join = [f for f in results if "run_sql(" in f]
    assert len(sheet_only) == 2
    assert len(join) == 1
    assert "session_duckdb()" in sheet_only[0]
    assert "scoped_dir" in join[0] and "zip_income.csv" in join[0]
    for formula in results:
        assert "data[1]" in formula
        assert SALES_NAMED_RANGE in formula or MARKETING_NAMED_RANGE in formula
        assert ".tolist()" not in formula
        assert "[.A" in formula  # explicit SQL cell/range, not magic-above
        assert len(formula) < 400
        for marker in _SQL_EMBED_MARKERS:
            assert marker not in formula, (marker, formula)
    # Live RESULTS formula cells must be unmerged (no number-columns-spanned).
    for m in re.finditer(r"<table:table-cell\b([^>]*)>", xml):
        attrs = m.group(1)
        if not _is_sql_results_formula(attrs):
            continue
        assert "number-columns-spanned" not in attrs, attrs


def _ods_row_is_empty(row: object) -> bool:
    from odf.table import TableCell
    from odf.text import P

    cells = row.getElementsByType(TableCell)  # type: ignore[attr-defined]
    if not cells:
        return True
    for cell in cells:
        if cell.getAttribute("formula"):
            return False
        texts = [str(p) for p in cell.getElementsByType(P) if str(p).strip()]
        if texts:
            return False
    return True


def _assert_ods_results_unmerged_with_clearance(sql_sheet: object) -> None:
    """Live ODS RESULTS cells have no column span; next 15 rows are empty."""
    from odf.table import TableCell, TableRow

    rows = list(sql_sheet.getElementsByType(TableRow))  # type: ignore[attr-defined]
    found = 0
    for idx, row in enumerate(rows):
        for cell in row.getElementsByType(TableCell):
            formula = cell.getAttribute("formula") or ""
            if not _is_sql_results_formula(formula):
                continue
            found += 1
            span = cell.getAttribute("numbercolumnsspanned")
            assert not span or int(span) <= 1, (idx, span, formula)
            empty = 0
            probe = idx + 1
            while probe < len(rows) and _ods_row_is_empty(rows[probe]):
                empty += 1
                probe += 1
            assert empty >= SQL_RESULTS_SPILL_GUTTER_ROWS, (idx, empty)
    assert found == 3


def test_build_ods_showcase_formulas_and_layout(tmp_path: Path) -> None:
    out_path = tmp_path / "python_showcase_demo.ods"
    build_ods_showcase(out_path)
    _assert_ods_formulas_and_layout(_ods_content_xml(out_path))


def test_shipped_ods_fixture_formulas_and_layout() -> None:
    _assert_ods_formulas_and_layout(_ods_content_xml(FIXTURE_ODS))
    from odf.opendocument import load
    from odf.table import Table

    doc = load(str(FIXTURE_ODS))
    sql_sheet = next(t for t in doc.spreadsheet.getElementsByType(Table) if t.getAttribute("name") == "SQL_DuckDB")
    _assert_ods_results_unmerged_with_clearance(sql_sheet)


def _assert_results_formula_is_short_and_quote_safe(
    formula: str, *, sql_range: str, join: bool = False
) -> str:
    """RESULTS =PY() is a short runner: SQL is a cell/range arg, not a giant string."""
    assert formula.startswith("="), formula
    payload = _quoted_formula_payload(formula)
    rest = formula[formula.rindex('"') + 1 :]
    assert len(payload) <= RESULTS_PY_CODE_MAX_LEN, (len(payload), payload)
    assert '"""' not in formula
    assert '"' not in payload
    assert "data[0]" in payload and "data[1]" in payload
    if join:
        assert "run_sql(" in payload
        assert "scoped_dir" in payload
        assert ZIP_INCOME_CSV_NAME in payload
        assert "session_duckdb" not in payload
    else:
        # only the table-name quotes in register('sales') / register('marketing')
        assert payload.count("'") <= 2
        assert "session_duckdb()" in payload
        assert "con.sql(sql)" in payload
        assert "result=con.sql(sql).df()" in payload
    assert ".tolist()" not in payload
    assert sql_range in rest
    assert SALES_NAMED_RANGE in rest or MARKETING_NAMED_RANGE in rest
    for marker in _SQL_EMBED_MARKERS:
        assert marker not in formula, marker
    # Premature string close made Calc treat SQL commas as OpenFormula
    # separators (Region, Category → region; category).
    assert "Region; Category" not in formula
    assert "SUM(Revenue);" not in formula
    return payload


def test_sql_results_gutter_covers_region_category_spill() -> None:
    """Sheet-only RESULTS need ≥15×5 empty (header + 12 Region×Category + clearance)."""
    assert SQL_RESULTS_SPILL_GUTTER_ROWS >= 15
    assert SQL_RESULTS_SPILL_GUTTER_COLS >= 5
    assert sql_results_gutter_rows("sheet_sales") == SQL_RESULTS_SPILL_GUTTER_ROWS
    assert sql_results_gutter_rows("sheet_marketing") == SQL_RESULTS_SPILL_GUTTER_ROWS
    assert sql_results_gutter_rows("join_zip") == SQL_RESULTS_SPILL_GUTTER_ROWS


def _empty_rows_below_xlsx_results(ws: object) -> list[tuple[str, int]]:
    """(RESULTS coordinate, consecutive empty rows in column A until the next value)."""
    from openpyxl.utils import coordinate_to_tuple

    gaps: list[tuple[str, int]] = []
    for coord, unused_formula in _xlsx_sql_duckdb_result_formulas(ws):
        row = coordinate_to_tuple(coord)[0]  # (row, column), 1-based
        empty = 0
        probe = row + 1
        while probe <= ws.max_row:  # type: ignore[attr-defined]
            val = ws[f"A{probe}"].value  # type: ignore[index]
            if val not in (None, ""):
                break
            empty += 1
            probe += 1
        gaps.append((coord, empty))
    return gaps


def _xlsx_merge_overlaps(ws: object, min_row: int, max_row: int, min_col: int, max_col: int) -> object | None:
    """Return a merged range that intersects the rectangle, else None."""
    for rng in ws.merged_cells.ranges:  # type: ignore[attr-defined]
        if rng.max_row < min_row or rng.min_row > max_row:
            continue
        if rng.max_col < min_col or rng.min_col > max_col:
            continue
        return rng
    return None


def _assert_xlsx_results_unmerged_with_clearance(ws: object) -> None:
    """Live RESULTS origin is unmerged; 15×5 under it (plus B–E of the origin row) is empty."""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import coordinate_to_tuple

    formulas = _xlsx_sql_duckdb_result_formulas(ws)
    assert len(formulas) == 3
    for coord, unused_formula in formulas:
        row, col = coordinate_to_tuple(coord)
        assert col == 1, coord
        origin_merge = _xlsx_merge_overlaps(ws, row, row, col, col)
        assert origin_merge is None, (coord, origin_merge)
        # Origin row B–E plus 15 rows × 5 cols under the formula.
        overlap = _xlsx_merge_overlaps(
            ws,
            row,
            row + SQL_RESULTS_SPILL_GUTTER_ROWS,
            1,
            SQL_RESULTS_SPILL_GUTTER_COLS,
        )
        assert overlap is None, (coord, overlap)
        for r in range(row, row + 1 + SQL_RESULTS_SPILL_GUTTER_ROWS):
            for c in range(1, SQL_RESULTS_SPILL_GUTTER_COLS + 1):
                if r == row and c == col:
                    continue
                cell = ws.cell(row=r, column=c)  # type: ignore[attr-defined]
                assert not isinstance(cell, MergedCell), cell.coordinate
                assert cell.value in (None, ""), (cell.coordinate, cell.value)


def test_generated_xlsx_results_have_spill_gutter(tmp_path: Path) -> None:
    """A 13-row spill from A19 must not hit the next section title."""
    out = tmp_path / "python_showcase_demo.xlsx"
    build_xlsx_showcase(out)
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb["SQL_DuckDB"]
    gaps = _empty_rows_below_xlsx_results(ws)
    assert len(gaps) == 3
    for coord, empty in gaps:
        assert empty >= SQL_RESULTS_SPILL_GUTTER_ROWS, (coord, empty)
    _assert_xlsx_results_unmerged_with_clearance(ws)
    assert SALES_NAMED_RANGE in wb.defined_names
    assert MARKETING_NAMED_RANGE in wb.defined_names
    assert wb.defined_names[SALES_NAMED_RANGE].attr_text == _xlsx_named_range_address(
        "Sales_Analytics", SALES_RANGE_XLSX
    )
    assert wb.defined_names[MARKETING_NAMED_RANGE].attr_text == _xlsx_named_range_address(
        "Statistics_ML", MARKETING_RANGE_XLSX
    )


def test_sheet_only_result_formulas_are_short_and_read_sql_from_cell_arg() -> None:
    """SQL is not embedded; RESULTS passes an explicit SQL cell/range plus the data range."""
    # Ranges here are formula-builder examples (scenario 1 SQL stays A11:A16).
    cases = (
        ("sheet_sales", "A11:A16", False),
        ("sheet_marketing", "A23:A29", False),
        ("join_zip", "A40:A54", True),
    )
    for kind, sql_range, is_join in cases:
        for ods_fmt in (False, True):
            formula = _scenario_result_formula(kind, sql_range, ods=ods_fmt)
            assert formula is not None, kind
            _assert_results_formula_is_short_and_quote_safe(
                formula, sql_range=sql_range, join=is_join
            )


def test_duckdb_sql_from_cell_code_is_quote_safe_and_under_cap() -> None:
    for table in ("sales", "marketing"):
        code = duckdb_sql_from_cell_code(table)
        assert len(code) <= RESULTS_PY_CODE_MAX_LEN
        assert '"' not in code
        assert f"register('{table}'" in code
        assert "session_duckdb()" in code
        assert "data[1]" in code
        assert "result=con.sql(sql).df()" in code
        assert ".tolist()" not in code


def test_duckdb_join_from_cell_code_is_quote_safe_and_under_cap() -> None:
    code = duckdb_join_from_cell_code("sales", ZIP_INCOME_CSV_NAME)
    assert len(code) <= RESULTS_PY_CODE_MAX_LEN
    assert '"' not in code
    assert "run_sql(" in code
    assert "scoped_dir" in code
    assert ZIP_INCOME_CSV_NAME in code
    assert "data[0]" in code and "data[1]" in code


def test_ods_formula_sql_results_two_args_not_rematched() -> None:
    formula = _scenario_result_formula("sheet_sales", "A11:A16", ods=True)
    assert formula is not None
    out = ods_formula(formula)
    # Named-range identity stays a name — ods_formula must not invent [$Sales_…].
    assert SALES_NAMED_RANGE in out
    assert "[$Sales_Analytics.A4:.J39]" not in out
    assert "[.A11:.A16]" in out
    assert _NESTED_SHEET_REF.search(out) is None
    assert out.startswith(f"of:={CALC_PYTHON_ADDIN_FN}(")
    for marker in _SQL_EMBED_MARKERS:
        assert marker not in out


def test_ods_formula_join_results_named_range_not_rematched() -> None:
    formula = _scenario_result_formula("join_zip", "A40:A54", ods=True)
    assert formula is not None
    out = ods_formula(formula)
    assert SALES_NAMED_RANGE in out
    assert "run_sql(" in out
    assert "[.A40:.A54]" in out
    assert _NESTED_SHEET_REF.search(out) is None


def test_sql_query_lines_keeps_visible_sql_out_of_the_formula() -> None:
    lines = sql_query_lines(SQL_SALES_BY_REGION_CATEGORY)
    assert any("SUM(Revenue)" in line for line in lines)
    assert len(lines) >= 4
    formula = _scenario_result_formula("sheet_sales", "A11:A16", ods=False)
    assert formula is not None
    assert "SUM(Revenue)" not in formula


def _xlsx_sql_duckdb_result_formulas(ws: object) -> list[tuple[str, str]]:
    found: list[tuple[str, str]] = []
    for row in ws.iter_rows():  # type: ignore[attr-defined]
        for cell in row:
            val = cell.value
            if isinstance(val, str) and val.startswith("=") and "data[1]" in val:
                found.append((cell.coordinate, val))
    return found


def test_fixture_xlsx_sql_results_formulas_are_short_and_quote_safe() -> None:
    """Committed showcase xlsx RESULTS cells stay short; SQL lives in cells."""
    from openpyxl import load_workbook

    path = Path(__file__).resolve().parents[1] / "fixtures" / "python_showcase_demo.xlsx"
    wb = load_workbook(path)
    ws = wb["SQL_DuckDB"]
    formulas = _xlsx_sql_duckdb_result_formulas(ws)
    assert len(formulas) == 3
    sql_text = "\n".join(
        str(cell.value) for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)
    )
    assert "SUM(Revenue)" in sql_text
    assert "SUM(Ad_Spend)" in sql_text
    assert "zip_income" in sql_text
    assert "{sheet" in sql_text and "{named_range" in sql_text
    assert SALES_NAMED_RANGE in wb.defined_names
    assert MARKETING_NAMED_RANGE in wb.defined_names
    assert wb.defined_names[SALES_NAMED_RANGE].attr_text == _xlsx_named_range_address(
        "Sales_Analytics", SALES_RANGE_XLSX
    )
    assert wb.defined_names[MARKETING_NAMED_RANGE].attr_text == _xlsx_named_range_address(
        "Statistics_ML", MARKETING_RANGE_XLSX
    )
    join_formulas = [f for unused_coord, f in formulas if "run_sql(" in f]
    sheet_formulas = [f for unused_coord, f in formulas if "con.sql(sql)" in f]
    assert len(join_formulas) == 1
    assert len(sheet_formulas) == 2
    for unused_coord, formula in formulas:
        # Trailing arg is the SQL cell/range on this sheet (sales stays A11:A16;
        # marketing sits below the RESULTS spill gutter).
        rest = formula[formula.rindex('"') + 1 :]
        sql_arg = rest.rsplit(",", 1)[-1].strip().rstrip(")")
        _assert_results_formula_is_short_and_quote_safe(
            formula, sql_range=sql_arg, join="run_sql(" in formula
        )
    gaps = _empty_rows_below_xlsx_results(ws)
    assert len(gaps) == 3
    for coord, empty in gaps:
        assert empty >= SQL_RESULTS_SPILL_GUTTER_ROWS, (coord, empty)
    _assert_xlsx_results_unmerged_with_clearance(ws)


def test_xlsx_relative_named_range_shift_matches_headed_binder_exceptions() -> None:
    """Repro path: relative Sheet!A4:J39 used from SQL_DuckDB RESULTS rows.

    UNO / Name Manager still show A4:J39 with Order_ID / Channel headers.
    Calc evaluates the name from the =PY() cell, so data[0] starts mid-table
    (scen1 A23 → row 26 Electronics/Consumer) or off the table (scen2 A51 →
    empty pack, Channel missing / candidate "column").
    """
    from plugin.scripting.calc_range import CalcRange

    sales = get_sales_dataset()
    marketing = get_marketing_dataset()
    sales_start = int("".join(ch for ch in SALES_RANGE_XLSX.split(":")[0] if ch.isdigit()))
    marketing_start = int("".join(ch for ch in MARKETING_RANGE_XLSX.split(":")[0] if ch.isdigit()))
    marketing_end = int("".join(ch for ch in MARKETING_RANGE_XLSX.split(":")[1] if ch.isdigit()))

    scen1_row = relative_named_range_eval_start_row(sales_start, 23)
    assert scen1_row == 26
    sales_offset = scen1_row - sales_start
    shifted_sales = sales[sales_offset:]
    assert shifted_sales[0][3] == "Electronics"
    assert shifted_sales[0][4] == "Consumer"
    defined_cols = list(CalcRange(sales).to_pandas().columns)
    shifted_cols = list(CalcRange(shifted_sales).to_pandas().columns)
    assert "Region" in defined_cols
    assert "Region" not in shifted_cols
    assert "Electronics" in shifted_cols and "Consumer" in shifted_cols

    scen2_row = relative_named_range_eval_start_row(marketing_start, 51)
    assert scen2_row == 54
    assert scen2_row > marketing_end
    marketing_offset = scen2_row - marketing_start
    shifted_marketing = marketing[marketing_offset:]
    assert shifted_marketing == []
    empty_cols = list(CalcRange(shifted_marketing).to_pandas().columns)
    assert "Channel" not in empty_cols


def test_xlsx_named_range_address_is_absolute() -> None:
    assert _xlsx_named_range_address("Sales_Analytics", SALES_RANGE_XLSX) == (
        "Sales_Analytics!$A$4:$J$39"
    )
    assert _xlsx_named_range_address("Statistics_ML", MARKETING_RANGE_XLSX) == (
        "Statistics_ML!$A$4:$G$24"
    )


def _ods_row_cell_texts(row: object) -> list[str]:
    from odf.table import TableCell
    from odf.text import P

    texts: list[str] = []
    for cell in row.getElementsByType(TableCell):  # type: ignore[attr-defined]
        for para in cell.getElementsByType(P):
            value = str(para).strip()
            if value:
                texts.append(value)
    return texts


def test_ods_and_xlsx_standard_sheets_share_row_skeleton(tmp_path: Path) -> None:
    """ODS standard sheets match the XLSX chrome: title, blank row 2, section 3, header 4."""
    from odf.opendocument import load
    from odf.table import Table, TableRow
    from openpyxl import load_workbook
    from openpyxl.utils import coordinate_to_tuple

    ods_path = tmp_path / "python_showcase_demo.ods"
    xlsx_path = tmp_path / "python_showcase_demo.xlsx"
    build_ods_showcase(ods_path)
    build_xlsx_showcase(xlsx_path)

    doc = load(str(ods_path))
    wb = load_workbook(xlsx_path)
    ods_tables = {str(t.getAttribute("name")): t for t in doc.spreadsheet.getElementsByType(Table)}
    assert list(ods_tables) == wb.sheetnames

    for spec in standard_sheet_specs():
        name = spec["name"]
        expected_title = f"📊 {name} — {spec['sub']}"
        header0 = spec["data"][0][0]
        ws = wb[name]
        assert ws["A1"].value == expected_title
        assert ws["A2"].value in (None, "")
        assert ws["A3"].value == spec["sec"]
        assert ws["A4"].value == header0

        rows = list(ods_tables[name].getElementsByType(TableRow))
        assert _ods_row_cell_texts(rows[0])[0] == expected_title
        assert _ods_row_cell_texts(rows[1]) == []
        assert _ods_row_cell_texts(rows[2])[0] == spec["sec"]
        assert _ods_row_cell_texts(rows[3])[0] == header0
        # Cell-less spacer serializes as <table:table-row …/> and Calc drops it.
        banner_texts = [_ods_row_cell_texts(row) for row in rows]
        assert any(STANDARD_METRICS_BANNER in texts for texts in banner_texts)

    xml = _ods_content_xml(ods_path)
    assert _ods_named_range_address("Sales_Analytics", SALES_RANGE_XLSX) in xml
    assert _ods_named_range_address("Statistics_ML", MARKETING_RANGE_XLSX) in xml
    assert "table:formula='" not in xml
    assert '<table:table-row table:style-name="row-spacer"/>' not in xml
    assert 'style:use-optimal-row-height="false"' in xml
    sa = _ods_table_xml(xml, "Sales_Analytics")
    assert re.search(
        r'table:style-name="row-spacer">\s*'
        r'<table:table-cell office:value-type="string">',
        sa,
    )
    assert re.search(
        r"row-section.*TRANSACTIONAL SALES DATASET.*row-header.*Order_ID",
        sa,
        flags=re.S,
    )

    xlsx_sql = wb["SQL_DuckDB"]
    xlsx_result_rows = [
        coordinate_to_tuple(coord)[0] for coord, unused_formula in _xlsx_sql_duckdb_result_formulas(xlsx_sql)
    ]
    ods_sql_rows = list(ods_tables["SQL_DuckDB"].getElementsByType(TableRow))
    ods_result_rows: list[int] = []
    for idx, row in enumerate(ods_sql_rows, start=1):
        from odf.table import TableCell

        for cell in row.getElementsByType(TableCell):
            formula = cell.getAttribute("formula") or ""
            if _is_sql_results_formula(formula):
                ods_result_rows.append(idx)
    assert ods_result_rows == xlsx_result_rows
    assert xlsx_sql.max_row == len(ods_sql_rows)


def test_generated_ods_formula_attrs_are_double_quoted(tmp_path: Path) -> None:
    out = tmp_path / "python_showcase_demo.ods"
    build_ods_showcase(out)
    xml = _ods_content_xml(out)
    assert "table:formula='" not in xml
    assert xml.count('table:formula="') >= 20


def _ods_table_xml(xml: str, name: str) -> str:
    start = xml.find(f'<table:table table:name="{name}"')
    assert start >= 0, name
    next_table = xml.find("<table:table table:name=", start + 1)
    return xml[start:next_table] if next_table >= 0 else xml[start:]


def _ods_row_slot_count(row_xml: str) -> int:
    """Logical columns: each table-cell or covered-table-cell is one slot."""
    return len(re.findall(r"<table:(?:table-cell|covered-table-cell)\b", row_xml))


def test_ods_overview_spans_emit_covered_cells_for_eight_columns(tmp_path: Path) -> None:
    """KPI/matrix rows must occupy A–H. Missing covered cells collapse later cards."""
    out = tmp_path / "python_showcase_demo.ods"
    build_ods_showcase(out)
    ov = _ods_table_xml(_ods_content_xml(out), "Overview")
    rows = re.findall(r"<table:table-row\b.*?</table:table-row>", ov, flags=re.S)
    labels = next(r for r in rows if "TOTAL REVENUE" in r)
    values = next(r for r in rows if "$119,142.00" in r)
    headers = next(r for r in rows if "Capability Domain" in r)
    hero = next(r for r in rows if "LibrePy / WriterAgent" in r)
    assert _ods_row_slot_count(hero) == 8
    assert _ods_row_slot_count(labels) == 8
    assert _ods_row_slot_count(values) == 8
    assert _ods_row_slot_count(headers) == 8
    assert labels.count("AVG PROFIT MARGIN") == 1
    assert labels.count("ANOMALIES FLAGGED") == 1
    assert labels.count("FORECAST TARGET") == 1
    assert headers.count("Traditional Calc Formula") == 1
    assert labels.count("<table:covered-table-cell") == 4
    assert headers.count("<table:covered-table-cell") == 4

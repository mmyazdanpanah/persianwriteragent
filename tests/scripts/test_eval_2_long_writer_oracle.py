# WriterAgent tests for scripts/eval_2_long_writer_oracle.py
from __future__ import annotations

import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from odf.opendocument import OpenDocumentText
from odf.text import P

_SCRIPTS = Path(__file__).resolve().parents[2] / "scripts"
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))

from eval_2_headed import main as headed_main  # noqa: E402
from eval_2_long_writer_oracle import (  # noqa: E402
    count_named_headings,
    detect_toc,
    main as oracle_main,
    read_comments,
    score_pack,
    score_text,
)

_PASSING = """
Northhaven Civic Library Capital Brief
Project code CL-2026-ANNEX

Purpose
Parks & Culture is briefing City Council on the Civic Library renovation
and annex in Northhaven.

Scope of Work
Base renovation of the Civic Library plus one annex option.

Budget and Schedule
Base renovation is $4.2 million. Annex option A is $1.8 million.
Construction window is April 2027 through October 2028.

Public-Access Impacts
Current Saturday hours stay posted until weekend staffing is decided.

Open Decisions
Annex siting, funding split, and weekend staffing still need a Council
decision.
"""

_OFFICE_NS = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
_TEXT_NS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"


def _padded() -> str:
    return _PASSING + (" Council review continues on the capital brief. " * 40)


def _write_docx(path: Path, text: str, *, headings: bool = False) -> Path:
    doc = Document()
    for line in text.strip().splitlines():
        if headings and line in {
            "Purpose",
            "Scope of Work",
            "Budget and Schedule",
            "Public-Access Impacts",
            "Open Decisions",
        }:
            doc.add_heading(line, level=1)
        else:
            doc.add_paragraph(line)
    doc.save(str(path))
    return path


def _write_odt_plain(path: Path, text: str) -> Path:
    doc = OpenDocumentText()
    for line in text.strip().splitlines():
        doc.text.addElement(P(text=line))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _inject_odt_structure(
    path: Path,
    *,
    toc: bool,
    headings: bool,
    comment: str | None,
) -> Path:
    with zipfile.ZipFile(path, "r") as zf:
        content = zf.read("content.xml").decode("utf-8")
        extras = {name: zf.read(name) for name in zf.namelist() if name != "content.xml"}
    root = ET.fromstring(content)
    body = root.find(f".//{{{_OFFICE_NS}}}text")
    assert body is not None
    if headings:
        for node in list(body):
            if node.tag == f"{{{_TEXT_NS}}}p" and (node.text or "").strip() in {
                "Purpose",
                "Scope of Work",
                "Budget and Schedule",
                "Public-Access Impacts",
                "Open Decisions",
            }:
                node.tag = f"{{{_TEXT_NS}}}h"
                node.set(f"{{{_TEXT_NS}}}style-name", "Heading_20_1")
                node.set(f"{{{_TEXT_NS}}}outline-level", "1")
    if toc:
        toc_el = ET.Element(f"{{{_TEXT_NS}}}table-of-content")
        toc_el.set(f"{{{_TEXT_NS}}}name", "TOC1")
        source = ET.SubElement(toc_el, f"{{{_TEXT_NS}}}table-of-content-source")
        source.set(f"{{{_TEXT_NS}}}outline-level", "3")
        index_body = ET.SubElement(toc_el, f"{{{_TEXT_NS}}}index-body")
        title = ET.SubElement(index_body, f"{{{_TEXT_NS}}}p")
        title.text = "Table of Contents"
        body.insert(0, toc_el)
    if comment:
        host = None
        for node in body.iter(f"{{{_TEXT_NS}}}p"):
            if "Annex siting" in "".join(node.itertext()):
                host = node
                break
        if host is None:
            host = ET.SubElement(body, f"{{{_TEXT_NS}}}p")
            host.text = "Annex siting remains open."
        ann = ET.SubElement(host, f"{{{_OFFICE_NS}}}annotation")
        ann.set(f"{{{_OFFICE_NS}}}name", "c1")
        para = ET.SubElement(ann, f"{{{_TEXT_NS}}}p")
        para.text = comment
    xml = ET.tostring(root, encoding="utf-8")
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in extras.items():
            compress = zipfile.ZIP_STORED if name == "mimetype" else zipfile.ZIP_DEFLATED
            zf.writestr(name, data, compress_type=compress)
        zf.writestr("content.xml", xml)
    return path


def _write_passing_odt(path: Path) -> Path:
    _write_odt_plain(path, _padded())
    return _inject_odt_structure(
        path,
        toc=True,
        headings=True,
        comment="Council still needs to pick annex option A vs B.",
    )


def _write_docx_with_toc_and_comment(path: Path) -> Path:
    _write_docx(path, _padded(), headings=True)
    with zipfile.ZipFile(path, "r") as zf:
        names = zf.namelist()
        extras = {name: zf.read(name) for name in names}
    # Minimal TOC field + comments part so the oracle can see both.
    document = extras["word/document.xml"].decode("utf-8")
    instr = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:instrText>TOC \\o &quot;1-3&quot;</w:instrText></w:r>"
    )
    document = document.replace("<w:body>", "<w:body>" + instr, 1)
    extras["word/document.xml"] = document.encode("utf-8")
    extras["word/comments.xml"] = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:comment w:id="0"><w:p><w:r><w:t>'
        "Council still needs to pick annex option A vs B."
        "</w:t></w:r></w:p></w:comment></w:comments>"
    ).encode("utf-8")
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in extras.items():
            zf.writestr(name, data)
    return path


def test_passing_odt_and_docx(tmp_path: Path) -> None:
    odt = _write_passing_odt(tmp_path / "ok.odt")
    result = score_pack(odt)
    assert result.passed, result.failures
    assert result.has_toc
    assert result.heading_count >= 3
    assert result.comment_count >= 1

    docx = _write_docx_with_toc_and_comment(tmp_path / "ok.docx")
    result = score_pack(docx)
    assert result.passed, result.failures


def test_empty_pack_fails(tmp_path: Path) -> None:
    path = _write_odt_plain(tmp_path / "empty.odt", "")
    result = score_pack(path)
    assert not result.passed
    assert any("empty" in item or "word_count" in item for item in result.failures)


def test_headings_without_toc_fail(tmp_path: Path) -> None:
    path = _write_odt_plain(tmp_path / "no-toc.odt", _padded())
    _inject_odt_structure(
        path,
        toc=False,
        headings=True,
        comment="Council still needs to pick annex option A vs B.",
    )
    result = score_pack(path)
    assert not result.passed
    assert any("table of contents" in item for item in result.failures)
    assert detect_toc(path) is False
    assert count_named_headings(path) >= 3


def test_toc_without_heading_styles_fails(tmp_path: Path) -> None:
    path = _write_odt_plain(tmp_path / "unstyled.odt", _padded())
    _inject_odt_structure(
        path,
        toc=True,
        headings=False,
        comment="Council still needs to pick annex option A vs B.",
    )
    result = score_pack(path)
    assert not result.passed
    assert any("heading styles" in item for item in result.failures)


def test_missing_comment_fails(tmp_path: Path) -> None:
    path = _write_odt_plain(tmp_path / "no-comment.odt", _padded())
    _inject_odt_structure(path, toc=True, headings=True, comment=None)
    result = score_pack(path)
    assert not result.passed
    assert any("review comment" in item for item in result.failures)
    assert read_comments(path) == []


def test_husk_comment_does_not_count(tmp_path: Path) -> None:
    path = _write_odt_plain(tmp_path / "husk-comment.odt", _padded())
    _inject_odt_structure(path, toc=True, headings=True, comment="Error: tool failed")
    result = score_pack(path)
    assert not result.passed
    assert any("review comment" in item or "husk" in item for item in result.failures)


def test_husk_body_fails() -> None:
    result = score_text(
        _padded() + "\nError: tool failed\n",
        para_count=16,
        heading_count=5,
        comment_count=1,
        has_toc=True,
    )
    assert not result.passed
    assert any("husk" in item for item in result.failures)


def test_missing_project_code_fails() -> None:
    text = _padded().replace("CL-2026-ANNEX", "XX-0000")
    result = score_text(
        text,
        para_count=16,
        heading_count=5,
        comment_count=1,
        has_toc=True,
    )
    assert not result.passed
    assert any("CL-2026-ANNEX" in item for item in result.failures)


def test_contents_index_equivalent_passes(tmp_path: Path) -> None:
    text = (
        "Table of Contents\n"
        "Purpose\n"
        "Scope of Work\n"
        "Budget and Schedule\n"
        "Public-Access Impacts\n"
        + _padded()
    )
    path = _write_odt_plain(tmp_path / "contents.odt", text)
    _inject_odt_structure(
        path,
        toc=False,
        headings=True,
        comment="Council still needs to pick annex option A vs B.",
    )
    result = score_pack(path)
    assert result.passed, result.failures
    assert result.has_toc


def test_headed_score_routes_to_long_writer_oracle(tmp_path: Path) -> None:
    pack = _write_passing_odt(tmp_path / "final_pack.odt")
    assert headed_main(["--task", "long-writer-pack", "--score", str(pack)]) == 0
    empty = _write_odt_plain(tmp_path / "empty.odt", "")
    assert headed_main(["--task", "long-writer-pack", "--score", str(empty)]) == 1


def test_oracle_cli_json(tmp_path: Path, capsys) -> None:
    pack = _write_passing_odt(tmp_path / "final_pack.odt")
    assert oracle_main(["--json", str(pack)]) == 0
    out = capsys.readouterr().out
    assert '"passed": true' in out
    assert '"has_toc": true' in out
